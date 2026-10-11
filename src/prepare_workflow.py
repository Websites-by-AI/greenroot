#!/usr/bin/env python3
from pathlib import Path
from collections import Counter
import json,importlib.util,csv
SRC=Path(__file__).resolve().parent
ROOT=SRC.parent/'site' if (SRC.parent/'site').is_dir() else SRC.parent
KIT=ROOT/'kit';KIT.mkdir(exist_ok=True)
def module(name,p):
 s=importlib.util.spec_from_file_location(name,p);m=importlib.util.module_from_spec(s);s.loader.exec_module(m);return m
c=module('workflow_content',SRC/'workflow-content.py');pc=module('partners_content',SRC/'partners-content.py')
d=json.loads((ROOT/'data/partners.json').read_text());wa=json.loads((SRC/'whatsapp-audit.json').read_text());routes=json.loads((SRC/'funding-routes.json').read_text())
# Use the exact, current primary programme list rather than the older search snippet.
routes['routes'][1]['source']='https://basvuru.itucekirdek.com/tr/basvurular'
routes['routes'][1]['apply']='https://basvuru.itucekirdek.com/discover/1/iklim_ve_temiz_teknolojiler_1'
(SRC/'funding-routes.json').write_text(json.dumps(routes,ensure_ascii=False,indent=2))
valid={r['id'] for r in d['contacts']}
for route in routes['routes']:
 assert all(x in valid for x in route.get('related',[])),route
 route['checked']=routes['checked'];route['phone_called']=False;route['whatsapp_status']='not_reviewed_for_this_route'
 for i,l in enumerate(['en','fa','tr']):
  profile='university' if route['kind'] in ['university','incubator'] else 'public_support'
  specific={'en':'For '+route['name']+', could you confirm the appropriate current route, prerequisite documents and whether an unincorporated team may apply? We will review the official source before any filing: ', 'fa':'برای '+(route.get('name_fa') or route['name'])+'، لطفاً مسیر فعلی مناسب، مدارک پیش‌نیاز و امکان درخواست تیم بدون شرکت را روشن کنید. پیش از هر ارسال، منبع رسمی را بررسی می‌کنیم: ', 'tr':route['name']+' için doğru güncel yolu, ön koşul belgelerini ve şirketleşmemiş ekibin başvuru imkanını teyit edebilir misiniz? Başvuru öncesi resmi kaynağı inceleyeceğiz: '}[l]
  route['message_subject_'+l]=c.SUBJECT[l].format(topic=c.TEXT[l]['ow.profile.'+profile])
  route['message_'+l]=c.INTRO[l].format(name=route.get('name_fa') if l=='fa' else route['name'])+c.KIND_PARAGRAPHS[profile][i]+'\n\n'+specific+route['source']+c.CLOSE[l]
for r in d['contacts']:
 r['whatsapp_status']=wa['records'][r['id']]['status']
 r['whatsapp_checked']=wa['checked'];r['whatsapp_channels']=wa['records'][r['id']]['channels']
 r['whatsapp_purpose']=wa['records'][r['id']]['purpose'];r['whatsapp_account_tested']=False
 r['whatsapp_source_pages']=wa['records'][r['id']]['source_pages']
 r['outreach_profile']='university' if r['country']=='TR' and any(w in r['parent'] for w in ['Üniversit','Teknik Üniversitesi']) else 'public_support' if r['name'].startswith('KOSGEB') else r['category'] if r['category'] in c.KIND_PARAGRAPHS else 'employment'
 r['fit_level']='conditional' if r['outreach_profile'] in ['university','public_support','research'] else 'high' if r['priority']==1 else 'medium'
 r['preferred_contact']='published_email' if r['email'] else 'official_contact_route'
 r['related_routes']=[x['id'] for x in routes['routes'] if r['id'] in x.get('related',[])]
 for i,l in enumerate(['en','fa','tr']):
  topic=c.TEXT[l]['ow.profile.'+r['outreach_profile']]
  r['message_subject_'+l]=c.SUBJECT[l].format(topic=topic)
  name=(r.get('name_fa') or r['name']) if l=='fa' else (r.get('name_en') or r['name'])
  specific=pc.OFFERS[r['offer']][l]
  marker={'en':'\n\nA specific topic for your team: ','fa':'\n\nموضوع مشخص برای تیم شما: ','tr':'\n\nEkibiniz için özel konu: '}[l]
  r['message_'+l]=c.INTRO[l].format(name=name)+c.KIND_PARAGRAPHS[r['outreach_profile']][i]+marker+specific+c.CLOSE[l]
counts=Counter(r['whatsapp_status'] for r in d['contacts'])
d['meta']['whatsapp_review']={'checked':wa['checked'],'pages_reviewed':wa['pages_reviewed'],'counts':dict(counts),'account_probing':False,'messages_sent':False}
d['meta']['workflow_profile']='Idea / no registered company; educational, residency, nationality and prior-support details unknown.'
(ROOT/'data/partners.json').write_text(json.dumps(d,ensure_ascii=False,indent=2))
(ROOT/'data/funding-routes.json').write_text(json.dumps(routes,ensure_ascii=False,indent=2))
js='/* Outreach workflow: no autonomous external actions. */\n(function(){var X='+json.dumps(c.TEXT,ensure_ascii=False)+';Object.keys(X).forEach(function(l){window.I18N[l]=Object.assign(window.I18N[l]||{},X[l]);});window.GR_WORKFLOW_CONFIG='+json.dumps({'api':'https://greenroot-followups.elasa2next.workers.dev','docs':c.DOCS,'defaultDelayDays':2,'sourceChecked':'2026-10-11'},ensure_ascii=False)+';})();\n'
(SRC/'workflow-i18n.js').write_text(js)
# Standalone private-workflow side instructions for a future connected assistant.
playbook='''# GreenRoot — Outreach-coach side instructions / دستور کناری مربی همکاری / İşbirliği koçu yönergeleri

## فارسی
- نقش: همراه آماده‌سازی همکاری و فاند؛ نه پزشک، وکیل، مدیر ثبت شرکت یا نماینده دارای اختیار امضا.
- وضعیت آغاز: ایده، بدون شرکت. تحصیل، اقامت، ملیت، سابقه حمایت، سهام‌داری و شواهد پژوهش‌وتوسعه نامعلوم‌اند؛ حدس نزن.
- هر ردیف بانک را جدا بخوان: نهاد مادر/شعبه/واحد، نقش پیشنهادی، شماره و منبع. احتمال پاسخ یا دریافت فاند را درصدسازی نکن.
- شماره موبایل مساوی واتساپ نیست. فقط لینک دارای شماره در صفحه رسمی را «منتشرشده» بنام؛ وجود حساب، فعالیت و اجازه پیام کاری را تأییدشده فرض نکن. کانال خدمات شهری/مشاوره را برای فروش استفاده نکن و کانال منتشرکننده فهرست را به مؤسسه دیگر نسبت نده.
- مخاطب و پیام را شخصی‌سازی کن: دانشگاه/TTO برای مربی پژوهش و پیش‌رشد؛ تکنوپارک برای پذیرش مناسب مرحله ایده؛ بازیافت برای مشخصات و خرید مواد؛ خیریه و درمان برای ارجاع و ایمنی؛ کسب‌وکار برای کار حقوق‌دار یا حمایت مشخص.
- دو مسیر گرین‌روت را حفظ کن: درآمد بازیافت برای هزینه درمان؛ و پژوهش مواد پزشکی فقط با تولید، کیفیت و مقررات دارای مجوز. اثر درمانی یا ایمنی محصول را جعل نکن.
- دانشگاه قبل از درخواست: معرفی تیم، فرضیه قابل‌آزمون، وضعیت واقعی نمونه اولیه، بازار، بودجه بسته کاری، ایمنی و اخلاق/حریم خصوصی. اطلاعات حساس و مدارک هویتی را در این سایت جمع نکن.
- ثبت شرکت با پذیرش پیش‌رشد، استقرار در تکنوپارک و دریافت سرمایه یکسان نیست. در BiGG ثبت زودهنگام شرکت/سهام‌داری می‌تواند مسیر را مخدوش کند؛ زمان شرط را از سند همان فراخوان بخوان.
- وضعیت تحصیل خارجی و مدارک YÖK، حمایت قبلی و سهام‌داری هنگام پذیرش مرحله اول را بررسی کن. این چک‌لیست تأیید رسمی صلاحیت نیست.
- 1812 سرمایه‌گذاری سهامی است؛ 1501/1507 گرنت پژوهشی در سطح شرکت؛ 1505 به پژوهشگر/دانشگاه مجری و شرکت مشتری نیاز دارد؛ حمایت توسعه کسب‌وکار KOSGEB را با گرنت بلاعوض اشتباه نگیر. قواعد و موعدها را پیش از پیشنهاد ارسال از منبع رسمی دوباره بخوان.
- پس از اعلام صریح کاربر که واقعاً ارسال کرده، موعد بررسی پاسخ را پیش‌فرض +۲ روز تقویمی بساز. کلیک روی ایمیل/واتساپ ارسال نیست.
- دو روز بعد اول صندوق پیام/پرتال را کاربر بررسی می‌کند؛ وانمود نکن به ایمیل او دسترسی داری. بی‌پاسخ: یک پیگیری مؤدبانه نزدیک روز ۷، با بازبینی و ارسال دستی. پاسخ مثبت: مواد خواسته‌شده و جلسه. رد/عدم‌تمایل: توقف.
- پس از پاسخ درباره شرایط، پرونده واقعی و فرم صحیح را آماده کن؛ پس از ارسال واقعی کاربر رسید/زمان را ثبت و موعد بررسی بساز. هیچ مرحله‌ای فقط به‌علت گذشت دو روز «پذیرش» یا «ثبت‌شده» نمی‌شود.
- قبل از ارسال، امضا، ثبت حقوقی، پرداخت، انتقال سهام یا اشتراک داده خصوصی تأیید مشخص کاربر و بررسی حرفه‌ای مناسب لازم است. ارسال انبوه و خودکار به نهادها ممنوع است.
- رها در سایت فعلی دموی قواعدمحور است. ادعای استدلال مستقل، خواندن صندوق ایمیل یا اجرای خودکار ثبت‌نام نکن.
- یادآور واقعی فقط با سرویس/فرستنده تنظیم‌شده، رضایت و آدرس تأییدشده کاربر. بدون آن صرفاً موعد محلی است. صف، پذیرش سرویس، تحویل و خواندن چهار وضعیت متفاوت‌اند.
- به سرور یادآور فقط آدرس مالک دفتر، نام عمومی نهاد، نوع کار و موعد برود؛ نه پیش‌نویس، یادداشت، رزومه، پرونده بیمار یا مدارک. توقف/حذف یادآور را همیشه فراهم کن.

## English
1. Coach preparation, not clinical care, legal registration, signing or financial commitments.
2. Start from a concept-stage, unincorporated team. Education, residence, nationality, previous support, shareholdings and demonstrated R&D are unknown; ask rather than invent.
3. Distinguish parent institutions, branches and units. Fit priority is editorial, never a fabricated probability of reply or funding.
4. WhatsApp needs a recipient-specific officially published link. No account probing; no inferred availability from mobile numbers. Do not inherit directory/HQ channels as local, or use service/counselling routes for a pitch.
5. Personalize by recipient role. Preserve both recycling-income-to-treatment and licensed, regulated medical-material R&D pathways. Never invent clinical or product-safety results.
6. Universities: concept note, honest team/skills, testable R&D, prototype evidence, market discovery, work packages, budget, ethics and privacy before a formal application. Do not collect ID or health records here.
7. Pre-incubation, company incorporation, technopark tenancy and investment are separate. Check the exact call's education, foreign-degree, prior-support and shareholding timing before recommending incorporation.
8. BiGG 1812 is equity investment, 1501/1507 company-level R&D grants, 1505 a joint executing research body/customer-company route, and KOSGEB business development includes repayable support. Recheck live official calls and gates.
9. Only after the user confirms an actual send, create a +2-calendar-day inbox/portal check. Link clicks are not sends. Do not claim inbox access.
10. No reply: consider ONE polite later follow-up around day 7, manually reviewed/sent. Positive: prepare requested materials/meeting. Decline: stop. Never turn elapsed time into registration or approval.
11. Log actual user submissions and receipts separately. Obtain specific permission before any external filing, signature, incorporation, payment, equity transfer or private-data disclosure; seek appropriate professional review. No bulk outreach.
12. Current Raha is a rules-based demo. Real email reminders require configured sender, consent and email verification. Queue, provider acceptance, delivery and reading are different states. Sync only minimal due-task data, not notes or drafts, and support deletion/unsubscribe.

## Türkçe
- Hazırlık koçusunuz; doktor, tescil/başvuru yetkilisi veya imza temsilcisi değilsiniz.
- Başlangıç fikir aşaması ve şirketsiz ekip. Eğitim, ikamet, vatandaşlık, önceki destek, ortaklık ve Ar-Ge kanıtını uydurmayın.
- Kurum/şube/birimi ayırın. Uyum önceliği yanıt veya fon olasılığı değildir.
- WhatsApp için resmi sayfada alıcı numaralı bağlantı gerekir. Hesap sorgusu yok; mobil numaradan varsayım yok. Rehber/merkez kanalı yerel şubeye atfedilmez; destek/danışma kanalında tanıtım yapılmaz.
- Mesajı üniversite/TTO, kuluçka, geri dönüşüm, bakım, belediye veya işverene göre uyarlayın. Tedaviyi iş çıktısına bağlamayın; tıbbi malzeme lisans/kalite/mevzuat gerektirir.
- Üniversite/teknopark öncesi: fikir notu, ekip, doğrulanabilir Ar-Ge hipotezi, prototip/THS kanıtı, pazar, iş paketleri, bütçe ve etik/mahremiyet. Hassas belge toplamayın.
- Ön kuluçka, şirket kuruluşu, teknopark yerleşimi ve yatırım ayrıdır. BiGG eğitim, YÖK belgeleri, önceki destek ve kabul tarihindeki ortaklık koşullarını güncel resmi çağrıdan okuyun; erken şirketleşmeyi otomatik önermeyin.
- 1812 pay yatırımı; 1501/1507 şirket Ar-Ge hibesi; 1505 araştırma kurumu + müşteri şirket; KOSGEB iş geliştirme geri ödemeli olabilir. Tarih ve uygunluğu yeniden teyit edin.
- Gerçek gönderimi kullanıcı açıkça kaydedince +2 takvim günü yanıt kontrolü oluşturun. Gelen kutusunu okuduğunuzu iddia etmeyin. Yanıt yoksa yaklaşık 7. günde tek nazik manuel takip; olumluysa belgeler/görüşme; ret varsa durun.
- Süre geçmesi kabul, kayıt veya yatırım değildir. Başvuru, imza, tescil, ödeme, pay devri ve özel veri paylaşımı öncesi açık onay ve uygun profesyonel inceleme alın. Toplu dış iletişim yok.
- Sitedeki Raha kural tabanlı demodur. Gerçek e-posta, yapılandırılmış hizmet/gönderen, onay ve doğrulanmış adres ister. Kuyruk, sağlayıcı kabulü, teslim ve okuma ayrıdır; yalnız asgari görev verisini eşleyin ve durdurma/silme sağlayın.
'''
for x in routes['routes']:
 playbook+='\n- '+x['name']+' — ['+str(x['ref'])+']('+x['source']+')'
(KIT/'GreenRoot-Raha-Outreach-Instructions-EN-FA-TR.md').write_text(playbook)
# Reviewable, organization-specific messages, not a bulk-send tool.
with (KIT/'GreenRoot-Personalized-Messages-210.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);cols=['id','name','outreach_profile','fit_level','whatsapp_status','whatsapp_checked','whatsapp_purpose','message_subject_en','message_en','message_subject_fa','message_fa','message_subject_tr','message_tr','source'];w.writerow(cols)
 for r in d['contacts']:w.writerow(["'"+str(r[k]) if str(r.get(k,'')).lstrip().startswith(('=','+','-','@')) else r.get(k,'') for k in cols])
# Editable documents. Unknown facts remain explicit placeholders; no fictional budgets/results.
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sections={
'en':[
 ('Status and purpose','DRAFT — concept-stage team; no registered company. Prepared for an eligibility / pre-incubation discussion, not proof of university admission, company registration, clinical outcomes, partner commitments or funding. Sources reviewed 11 October 2026.'),
 ('1. Project identity and ask','Project: GreenRoot. Applicant / team: [fill]. Country of residence / intended incorporation: [fill]. University affiliation and degree evidence: [unknown — verify]. Target programme / current call / deadline: [fill from official source]. Requested meeting or assistance: [one specific ask].'),
 ('2. Problem, user and payer','Describe the adult target group, work barriers and waste-flow problem without identifying individuals. Distinguish the participant, customer, recycler/offtaker and sponsor/payer. Validate the need with interviews only after informed consent; do not fabricate quotations.'),
 ('3. Proposed GreenRoot model','Voluntary lawful paid green work; independent licensed treatment; recycling revenue may support treatment costs. A separate, research-led material-to-medical-supplies pathway needs licensed manufacture, material/quality testing and regulatory clearance. No therapeutic efficacy or medical-product safety is claimed.'),
 ('4. Technology / R&D hypothesis','Choose and substantiate one genuine technology question: e.g. responsible multilingual employment navigation, privacy-preserving outcome/flow tracking, or licensed circular-material quality research. State the technical uncertainty, novelty versus existing products, test method and success/failure criteria. A static site or social mission alone is not R&D evidence. TRL/THS: [unknown; attach actual evidence].'),
 ('5. Prototype and validation','What exists now: a concept website, a rules-based Raha demo and an institutional-source outreach bank. Record additional real prototypes/results separately. Existing sample impact figures are demo data and must not be submitted as achieved outcomes. Evidence / attachments: [fill].'),
 ('6. Team and academic route','Founder roles, relevant skills, time commitment and missing competencies: [fill]. Proposed mentor/lab/clinical/manufacturing partner: [not agreed unless written evidence]. Student/graduation status, foreign-degree YÖK recognition/equivalence, residence/work permissions: [verify with the actual programme]. Do not attach IDs to the public website.'),
 ('7. Work packages and timetable','Prepare WP1 needs/market validation; WP2 prototype/technical test; WP3 safety and independent ethical/regulatory review; WP4 small pilot and evaluation; WP5 commercial/operational validation. Dates and owners: [fill]. Clinical studies require qualified oversight and applicable ethics/clinical approvals, not a website checkbox.'),
 ('8. Budget and financing route','All costs below are blanks to be quoted, not estimates presented as facts. Identify own contribution, cash-flow timing, approved eligible costs and funding gap. Distinguish grant, equity and repayable support. Do not borrow, form a company or transfer shares merely to meet a deadline.'),
 ('9. Safety, ethics and privacy','Adults only for proposed work; genuine wages/contracts/insurance; non-hazardous waste and professional safety supervision; care independent of work output; appropriate clinical licensing; minimal data and informed consent; separate health/work records. No patient records, identity documents or keys in the public repository or reminder service.'),
 ('10. Application gates','Pre-incubation ≠ technopark tenancy ≠ company incorporation ≠ investment. Check current intake, technology criteria, fees, academic status, previous support, shareholding timing and foreign-founder requirements. For BiGG follow the current agency/PRODİS process; for company grants first verify the legal/company gates. Record actual filing receipts only after submission.'),
 ('11. Contacts and next actions','Choose one route and one named institutional office from the published sources. Edit a recipient-specific introduction. After YOU actually send it, record a +2-day inbox/portal check; if unanswered, consider one later polite follow-up. Stop after a decline. Real email reminders require a configured sender and a verified address.'),
 ('12. Approval before submission','Applicant review: [not completed]. Programme eligibility confirmed by: [not confirmed]. Professional legal/financial/clinical review where needed: [not completed]. Signatures, incorporation, fees and equity agreements: [not authorized].')],
'fa':[
 ('وضعیت و هدف','پیش‌نویس — تیم در مرحله ایده، بدون شرکت ثبت‌شده. برای بررسی شرایط و پیش‌رشد آماده شده؛ مدرک پذیرش دانشگاه، ثبت شرکت، نتیجه بالینی، تعهد شریک یا دریافت فاند نیست. بازبینی منابع: ۱۱ اکتبر ۲۰۲۶.'),
 ('۱. هویت پروژه و درخواست','پروژه: گرین‌روت. متقاضی و تیم: [تکمیل شود]. کشور اقامت و ثبت احتمالی: [تکمیل شود]. وابستگی دانشگاهی و مدرک تحصیل: [نامعلوم؛ بررسی شود]. برنامه، فراخوان جاری و موعد: [از منبع رسمی]. درخواست جلسه یا کمک: [یک درخواست روشن].'),
 ('۲. مسئله، کاربر و پرداخت‌کننده','گروه هدف بزرگسال، موانع کار و مسئله جریان پسماند را بدون شناسایی افراد توضیح دهید. شرکت‌کننده، مشتری، خریدار مواد و حامی/پرداخت‌کننده جدا باشند. مصاحبه نیازسنجی با رضایت آگاهانه؛ بدون نقل‌قول ساختگی.'),
 ('۳. مدل پیشنهادی گرین‌روت','کار سبز داوطلبانه، قانونی و حقوق‌دار؛ درمان مستقل و دارای مجوز؛ درآمد بازیافت می‌تواند هزینه درمان را پشتیبانی کند. مسیر جداگانه پژوهش مواد تا لوازم پزشکی، نیازمند تولید دارای مجوز، آزمون مواد و کیفیت و تأیید مقررات است. اثر درمانی یا ایمنی محصول پزشکی ادعا نمی‌شود.'),
 ('۴. فرضیه فناوری و پژوهش‌وتوسعه','یک سؤال فناوری واقعی انتخاب و مستند کنید: راهبری شغلی چندزبانه مسئولانه، ردیابی جریان/پیامد با حفاظت داده یا پژوهش کیفیت مواد چرخشی با شریک مجاز. ابهام فنی، نوآوری نسبت به محصولات موجود، روش آزمون و معیار موفقیت/شکست را بنویسید. سایت ثابت یا مأموریت اجتماعی به‌تنهایی مدرک پژوهش نیست. آمادگی فناوری: [نامعلوم؛ مدرک واقعی].'),
 ('۵. نمونه اولیه و اعتبارسنجی','آنچه اکنون وجود دارد: سایت مفهومی، دموی قواعدمحور رها و بانک ارتباط با منابع مؤسسه‌ای. نمونه یا نتیجه واقعی دیگر را جدا ثبت کنید. آمار نمونه اثر در سایت دمو است و نباید نتیجه محقق‌شده در درخواست فاند معرفی شود. شواهد و پیوست‌ها: [تکمیل شود].'),
 ('۶. تیم و مسیر دانشگاهی','نقش، مهارت، تعهد زمانی و کمبود توانمندی: [تکمیل شود]. مربی، آزمایشگاه، شریک درمانی یا تولیدی: [توافق نشده مگر با مدرک]. وضعیت تحصیل، شناسایی/معادل‌سازی تحصیل خارجی YÖK، اقامت و اجازه کار: [با برنامه واقعی بررسی شود]. مدارک هویتی در سایت عمومی بارگذاری نشود.'),
 ('۷. بسته‌های کاری و زمان‌بندی','بسته ۱ نیاز و بازار؛ ۲ نمونه و آزمون فنی؛ ۳ ایمنی و بررسی مستقل اخلاق و مقررات؛ ۴ پایلوت کوچک و ارزیابی؛ ۵ اعتبارسنجی تجاری و عملیاتی. تاریخ و مسئول: [تکمیل شود]. پژوهش بالینی به نظارت متخصص و مجوز لازم نیاز دارد، نه تیک سایت.'),
 ('۸. بودجه و مسیر منابع','هزینه‌ها باید با استعلام تکمیل شوند؛ عدد نمونه، واقعیت مالی نیست. سهم خودتان، جریان نقدی، هزینه واجدشرایط و کسری بودجه جدا مشخص شود. گرنت، سرمایه سهامی و حمایت بازپرداختی را تفکیک کنید. برای رسیدن به موعد، بدون بررسی بدهی، شرکت یا انتقال سهام ایجاد نکنید.'),
 ('۹. ایمنی، اخلاق و حریم خصوصی','کار پیشنهادی برای بزرگسال؛ قرارداد، دستمزد و بیمه واقعی؛ پسماند غیرخطرناک و نظارت حرفه‌ای؛ مراقبت مستقل از خروجی کار؛ مجوز بالینی؛ رضایت و حداقل داده؛ جداسازی داده سلامت و کار. پرونده بیمار، مدارک هویتی و کلیدها در مخزن عمومی یا سرویس یادآور قرار نگیرند.'),
 ('۱۰. شرایط درخواست','پیش‌رشد، استقرار تکنوپارک، ثبت شرکت و سرمایه‌گذاری یکسان نیستند. پذیرش فعلی، موضوع فناوری، هزینه، تحصیل، حمایت قبلی، زمان شرط سهام‌داری و قواعد بنیان‌گذار خارجی را بررسی کنید. BiGG طبق نهاد مجری و PRODİS جاری؛ گرنت شرکت پس از تأیید شرایط شرکت. رسید فقط بعد از ارسال واقعی.'),
 ('۱۱. تماس و اقدام بعدی','یک مسیر و دفتر مؤسسه‌ای واقعی از منابع انتخاب کنید. معرفی مخصوص مخاطب را ویرایش کنید. پس از ارسال واقعی خودتان، بررسی پاسخ دو روز بعد ثبت شود؛ بی‌پاسخ: یک پیگیری محترمانه دیرتر. رد: توقف. یادآور واقعی به فرستنده تنظیم‌شده و آدرس تأییدشده نیاز دارد.'),
 ('۱۲. تأیید قبل از ارسال','بازبینی متقاضی: [انجام نشده]. شرایط برنامه توسط: [تأیید نشده]. بررسی حقوقی، مالی و بالینی لازم: [انجام نشده]. امضا، ثبت شرکت، هزینه و توافق سهام: [مجاز نشده].')],
 'tr':[
 ('Durum ve amaç','TASLAK — fikir aşamasındaki ekip, kayıtlı şirket yok. Uygunluk / ön kuluçka görüşmesi için hazırlanmıştır; kabul, tescil, klinik sonuç, ortak taahhüdü veya fon kanıtı değildir. Kaynak incelemesi: 11 Ekim 2026.'),
 ('1. Proje ve talep','Proje: GreenRoot. Başvuran/ekip: [doldurun]. İkamet ve planlanan kuruluş ülkesi: [doldurun]. Üniversite bağı/eğitim belgesi: [bilinmiyor; teyit]. Program, güncel çağrı ve tarih: [resmi kaynaktan]. Tek açık görüşme talebi: [doldurun].'),
 ('2. Sorun, kullanıcı ve ödeyen','Yetişkin hedef grup, iş engelleri ve atık akışını kişileri tanımlamadan açıklayın. Katılımcı, müşteri, alıcı ve sponsor/ödeyeni ayırın. İhtiyaç görüşmeleri bilgilendirilmiş onayla yapılır; uydurma alıntı yok.'),
 ('3. Önerilen GreenRoot modeli','Gönüllü yasal ücretli yeşil iş; bağımsız lisanslı tedavi; geri dönüşüm geliri tedavi giderlerini destekleyebilir. Ayrı tıbbi malzeme araştırma yolu lisanslı üretim, malzeme/kalite testi ve mevzuat onayı gerektirir. Tedavi etkisi veya tıbbi ürün güvenliği iddiası yok.'),
 ('4. Teknoloji / Ar-Ge hipotezi','Gerçek bir teknik soru seçin: sorumlu çok dilli istihdam yönlendirmesi, mahremiyet koruyan akış/sonuç takibi veya lisanslı döngüsel malzeme kalitesi araştırması. Teknik belirsizlik, mevcut ürüne göre yenilik, test ve başarı/başarısızlık ölçütlerini yazın. Statik site veya sosyal amaç tek başına Ar-Ge kanıtı değildir. THS: [bilinmiyor; gerçek kanıt].'),
 ('5. Prototip ve doğrulama','Mevcut: fikir sitesi, kural tabanlı Raha demosu ve kurumsal kaynaklı iletişim rehberi. Diğer gerçek prototip/sonuçları ayrı kaydedin. Örnek etki verileri demodur; gerçekleşmiş sonuç diye sunulmaz. Kanıt/ekler: [doldurun].'),
 ('6. Ekip ve akademik yol','Rol, beceri, zaman ve eksikler: [doldurun]. Mentor/laboratuvar/klinik/üretim ortağı: [yazılı kanıt yoksa anlaşılmadı]. Eğitim, yabancı eğitim YÖK tanıma/denklik, ikamet ve çalışma izinleri: [gerçek programla teyit]. Kimlik belgesi herkese açık siteye yüklenmez.'),
 ('7. İş paketleri ve takvim','İP1 ihtiyaç/pazar; İP2 prototip/teknik test; İP3 güvenlik ve bağımsız etik/mevzuat; İP4 küçük pilot/değerlendirme; İP5 ticari/operasyonel doğrulama. Tarih ve sorumlu: [doldurun]. Klinik çalışma, site onayı değil yetkin gözetim ve gereken izinleri ister.'),
 ('8. Bütçe ve finansman','Maliyetler tekliflerle doldurulacak boşluklardır; örnek sayı mali gerçek değildir. Öz katkı, nakit zamanlaması, uygun gider ve fon açığını ayırın. Hibe, pay yatırımı ve geri ödeme farklıdır. Sırf süreye yetişmek için borç, şirket veya pay devri yapmayın.'),
 ('9. Güvenlik, etik ve mahremiyet','Yetişkin iş planı; gerçek ücret/sözleşme/sigorta; tehlikesiz atık ve profesyonel güvenlik; bakım iş çıktısından bağımsız; klinik lisans; asgari veri ve onay; sağlık/iş kayıtları ayrı. Hasta/kimlik kayıtları veya anahtarlar kamu deposu/hatırlatma hizmetine konmaz.'),
 ('10. Başvuru koşulları','Ön kuluçka, teknopark yerleşimi, tescil ve yatırım ayrıdır. Güncel kabul, teknoloji, ücret, eğitim, önceki destek, ortaklık zamanı ve yabancı kurucu koşullarını kontrol edin. BiGG için güncel kuruluş/PRODİS; şirket hibeleri için şirket koşulları. Makbuz yalnız gerçek gönderimden sonra.'),
 ('11. İletişim ve sonraki adım','Resmi kaynaktan bir yol ve kurumsal ofis seçin. Alıcıya özel tanıtımı düzenleyin. Gerçek gönderiminizden sonra +2 gün yanıt kontrolü; yanıt yoksa bir sonraki nazik takip; ret varsa durun. Gerçek e-posta için yapılandırılmış gönderen ve doğrulanmış adres gerekir.'),
 ('12. Gönderim öncesi onay','Başvuran incelemesi: [yapılmadı]. Program uygunluğu: [teyit edilmedi]. Gereken hukuki, mali ve klinik inceleme: [yapılmadı]. İmza, tescil, ücret veya pay sözleşmesi: [yetki verilmedi].')]
}
for l in ['en','fa','tr']:
 doc=Document();sec=doc.sections[0];sec.header.paragraphs[0].text='GREENROOT / APPLICATION PREPARATION · DRAFT'
 doc.add_heading('GreenRoot — '+{'en':'Candidate dossier','fa':'پرونده پیشنهادی درخواست','tr':'Aday başvuru dosyası'}[l],0)
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(9)
 for title,text in sections[l]:doc.add_heading(title,level=1);doc.add_paragraph(text)
 table=doc.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
 for cell,txt in zip(table.rows[0].cells,{'en':['Work package / cost','Quote basis','Amount / currency','Eligible? / own contribution'],'fa':['بسته کاری / هزینه','مبنای استعلام','مبلغ / ارز','واجدشرایط؟ / سهم خودتان'],'tr':['İş paketi / maliyet','Teklif temeli','Tutar / para birimi','Uygunluk / öz katkı']}[l]):cell.text=txt
 for label in {'en':['Research / prototype','Safety / insurance','Licensed partner / quality','Personnel / real wages','Legal / application fee'],'fa':['پژوهش / نمونه اولیه','ایمنی / بیمه','شریک دارای مجوز / کیفیت','نیروی کار / دستمزد واقعی','حقوقی / هزینه درخواست'],'tr':['Araştırma / prototip','Güvenlik / sigorta','Lisanslı ortak / kalite','Personel / gerçek ücret','Hukuki / başvuru gideri']}[l]:
  cells=table.add_row().cells;cells[0].text=label
  for cell in cells[1:]:cell.text={'en':'[fill]','fa':'[تکمیل شود]','tr':'[doldurun]'}[l]
 doc.add_heading({'en':'Source-backed routes (verify current rules)','fa':'مسیرهای منبع‌دار — قواعد فعلی تأیید شود','tr':'Kaynaklı yollar — güncel kuralları teyit edin'}[l],level=1)
 for x in routes['routes']:doc.add_paragraph(x['name']+'\n'+x['summary_'+l]+'\n'+x['source'])
 if l=='fa':
  allparas=list(doc.paragraphs)+[p for tbl in doc.tables for row in tbl.rows for cell in row.cells for p in cell.paragraphs]
  for p in allparas:
   prop=p._p.get_or_add_pPr();bidi=OxmlElement('w:bidi');bidi.set(qn('w:val'),'1');prop.append(bidi)
 doc.save(KIT/f'GreenRoot-Application-Dossier-{l.upper()}.docx')
print('Prepared 210 personalized messages, 13 source-backed routes, 3 editable dossiers and Raha instructions; WhatsApp',dict(counts))
